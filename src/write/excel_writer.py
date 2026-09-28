"""JSON -> Excel 작성기.

역할:
- 통합 템플릿에 값만 써 넣음 (라벨/용어 정규화는 src.normalize에 있음)
- 공개 진입점:
  - json_to_excel(...)
  - process_all_json(json_dir, excel_dir, survey_type, org_name)
"""

from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from ..config.org import load_org_config
from ..util.cancel import UserCancelled
from ..util.paths import project_root as _project_root
from ..util.unicode_name import maybe_decode_hashu
from ..normalize import normalize_table_json
from ..normalize.regions import extract_region_list_for_table, is_overall_region_name, normalize_region_items
from .layouts import find_category_row
from .table_finalize import finalize_table_styles


def _merged_end_row(ws, row: int, col: int) -> int:
    """(row, col)이 속한 병합 영역의 max_row를 반환(없으면 row)."""
    coord = ws.cell(row=row, column=col).coordinate
    for mr in ws.merged_cells.ranges:
        try:
            if coord in mr:
                return int(mr.max_row)
        except Exception:
            continue
    return row


def _is_merged(ws, row: int, col: int) -> bool:
    return isinstance(ws.cell(row=row, column=col), MergedCell)


def get_unique_filepath(output_dir: str, filename: str) -> str:
    """파일이 존재하면 (1), (2)...를 붙여 고유 경로 생성."""
    base, ext = os.path.splitext(filename)
    path = os.path.join(output_dir, filename)
    if not os.path.exists(path):
        return path
    i = 1
    while True:
        new_path = os.path.join(output_dir, f"{base} ({i}){ext}")
        if not os.path.exists(new_path):
            return new_path
        i += 1


def _parse_survey_period(period_str: str) -> tuple[Optional[int], Optional[int]]:
    import re

    year_match = re.search(r"(\d{4})년", period_str)
    month_match = re.search(r"(\d{1,2})월", period_str)
    year = int(year_match.group(1)) if year_match else None
    month = int(month_match.group(1)) if month_match else None
    return year, month


def _parse_survey_method(raw_method: str) -> tuple[str, str]:
    """조사방법을 템플릿용 method/frame으로 분리."""
    if not raw_method:
        return "", ""

    # 휴리스틱: '/' 또는 ' / ' 기준으로 분리
    parts = [p.strip() for p in raw_method.split("/") if p.strip()]
    if len(parts) == 1:
        return parts[0], ""
    if len(parts) >= 2:
        return parts[0], "/".join(parts[1:])
    return raw_method, ""


def write_metadata(ws, metadata: dict, title: str = None, survey_type: str = "전국"):
    """통합 템플릿의 고정 셀에 메타데이터 기록."""

    # E2: 조사 기관
    if "조사기관" in metadata:
        ws.cell(row=2, column=5).value = metadata["조사기관"]

    # G2: 조사 기간
    if "조사기간" in metadata:
        ws.cell(row=2, column=7).value = metadata["조사기간"]

        if survey_type == "전국":
            year, month = _parse_survey_period(str(metadata["조사기간"]))
            if year:
                ws.cell(row=2, column=14).value = year
            if month:
                ws.cell(row=2, column=15).value = month

    # 지방: N2=지역
    if survey_type == "지방" and "지역" in metadata:
        ws.cell(row=2, column=14).value = metadata["지역"]

    # H2/I2: 조사 방법/추출틀
    if "조사방법" in metadata:
        method, frame = _parse_survey_method(str(metadata["조사방법"]))
        ws.cell(row=2, column=8).value = method
        ws.cell(row=2, column=9).value = frame

    # J2: 표본크기
    if "표본크기" in metadata:
        ws.cell(row=2, column=10).value = metadata["표본크기"]

    # K2: 응답률
    if "응답률" in metadata:
        ws.cell(row=2, column=11).value = metadata["응답률"]

    # L2: 표본오차
    if "표본오차" in metadata:
        ws.cell(row=2, column=12).value = metadata["표본오차"]

    # M2: 타이틀
    if title:
        ws.cell(row=2, column=13).value = title


def write_response_headers(ws, response_items: List[str], header_row: int = 4, start_col: int = 5):
    for i, item in enumerate(response_items):
        ws.cell(row=header_row, column=start_col + i).value = item



def _read_template_response_headers(ws, header_row: int = 4, start_col: int = 5, max_cols: int = 60) -> List[str]:
    """템플릿에서 응답 헤더 라벨을 읽어옴(고정 헤더 보존용).

    PSR처럼 고정 응답 템플릿은 PDF의 동적 라벨로 덮어쓰지 않는다.
    """

    items: List[str] = []
    for i in range(max_cols):
        v = ws.cell(row=header_row, column=start_col + i).value
        if v is None:
            break
        s = str(v).strip()
        if not s:
            break
        items.append(s)
    return items


def write_data_row(ws, row_num: int, values: dict, response_items: List[str], start_col: int = 4):
    if not row_num:
        return

    # D열: 조사완료
    if isinstance(values, dict) and "조사완료" in values:
        ws.cell(row=row_num, column=start_col).value = values["조사완료"]

    for i, item in enumerate(response_items):
        v = values.get(item) if isinstance(values, dict) else None
        cell = ws.cell(row=row_num, column=start_col + 1 + i)
        cell.value = v if v is not None else None


def write_region_info(ws, region_info, region_count: int):
    if not region_info:
        return

    # 지방 템플릿은 '지역별' 라벨이 B14:B17처럼 병합되어 있어
    # 지역 수(region_count)가 적으면 기존 로직(14+region_count+1)이 병합 셀 내부를 가리켜 오류가 납니다.
    # - 마지막 지역 행(last_region_row) 기준 1줄 띄우고
    # - 병합 구간 끝(merge_end)보다 아래에서 시작하도록 안전한 시작 행을 계산합니다.
    merge_end = _merged_end_row(ws, row=14, col=2)  # B14
    last_region_row = 14 + max(int(region_count or 0), 1) - 1
    start_row = max(last_region_row + 2, merge_end + 1)

    regions = normalize_region_items(region_info)

    # 권역설명에 "전체" 행이 붙는 경우가 있어, 다른 권역이 있으면 제외해 일관성 유지
    non_overall = [r for r in regions if not is_overall_region_name(r.get("name", ""))]
    if non_overall:
        regions = non_overall

    # 템플릿 수정으로 병합 셀이 생겼을 수 있으니 병합 영역을 건너뛰며 기록
    def _next_unmerged(r: int, c: int) -> int:
        rr = r
        while isinstance(ws.cell(row=rr, column=c), MergedCell):
            rr += 1
        return rr

    row = start_row
    for rinfo in regions:
        row = _next_unmerged(row, 2)
        ws.cell(row=row, column=2).value = rinfo.get("name", "")
        ws.cell(row=row, column=3).value = rinfo.get("areas", "")
        row += 1


def _get_template_path(project_root: Path, survey_type: str, item_type: str) -> Path:
    folder = "regular_poll" if survey_type == "전국" else "local_elections"
    return project_root / "templates" / folder / f"template{item_type}.xlsx"


def json_to_excel(
    json_path: str,
    excel_path: str,
    survey_type: str,
    org_name: str,
    metadata: dict | None = None,
    region_info: Any | None = None,
) -> Dict[str, Any]:
    """단일 표 JSON을 Excel 파일로 변환."""

    project_root = _project_root()
    org = load_org_config(org_name, project_root=project_root)

    with open(json_path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    item_type = (raw.get("항목유형") or "ISSUE").upper()
    item_name = raw.get("항목명") or raw.get("질문") or Path(json_path).stem

    template_path = _get_template_path(project_root, survey_type, item_type)
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    wb = load_workbook(template_path)
    ws = wb.active

    # 제목을 B2에 표시
    ws.cell(row=2, column=2).value = item_name

    # 항목 간 공유되는 전달 메타데이터를 우선 사용
    meta_use = metadata if metadata is not None else raw.get("메타데이터", {})
    write_metadata(ws, meta_use or {}, title=item_name, survey_type=survey_type)

    # 쓰기 전에 표를 정규화
    norm, norm_report = normalize_table_json(raw, org, survey_type, region_info=region_info)

    response_items = norm.get("응답항목") or []
    table_data = norm.get("데이터") or {}

    # ISSUE: 응답 항목이 매번 달라 헤더를 작성
    if item_type == "ISSUE" and isinstance(response_items, list):
        write_response_headers(ws, response_items)

    # PSR: 응답 열이 고정이므로 템플릿 헤더를 그대로 사용
    if item_type == "PSR":
        response_items = _read_template_response_headers(ws)

    # 지방은 동적 지역 행 라벨을 C열에 직접 기록
    region_list: List[str] = []
    if survey_type == "지방":
        region_info_use = region_info if region_info is not None else norm.get("권역설명")
        # 표에 실제 행 라벨(구/군/시 등) 목록이 있으면 그 목록을 우선 사용
        table_cats = list((table_data or {}).keys()) if isinstance(table_data, dict) else []
        region_list = extract_region_list_for_table(region_info_use, table_cats)

    not_found = []
    written = 0
    for cat, values in (table_data or {}).items():
        row_num = find_category_row(cat, survey_type, region_list)
        if row_num is None:
            not_found.append(cat)
            continue
        if survey_type == "지방" and row_num >= 14:
            ws.cell(row=row_num, column=3).value = cat
        write_data_row(ws, row_num, values, response_items)
        written += 1

    # 권역설명 블록(지방 전용)
    if survey_type == "지방":
        region_info_use = region_info if region_info is not None else norm.get("권역설명")
        if region_info_use:
            write_region_info(ws, region_info_use, len(region_list))

    # 필요 시 기관별 Excel 패치 적용
    patches = org.excel_patches or {}
    side_key = "전국" if survey_type == "전국" else "지방"
    side_patch = patches.get(side_key) if isinstance(patches.get(side_key), dict) else {}
    # 지원 항목: set_cells: [{row:31,col:3,value:"..."}, ...]
    if isinstance(side_patch, dict) and isinstance(side_patch.get("set_cells"), list):
        for c in side_patch.get("set_cells"):
            if not isinstance(c, dict):
                continue
            r = c.get("row")
            col = c.get("col")
            val = c.get("value")
            try:
                r_i = int(r)
                c_i = int(col)
            except Exception:
                continue
            ws.cell(row=r_i, column=c_i).value = val

    # 표 스타일 마무리(동적 테두리/병합/형식)
    # - 경우 A: 데이터가 템플릿 영역을 넘으면 오른쪽/아래로 확장
    # - 경우 B: 권역설명 블록은 건드리지 않고 표 영역만 수정
    # - 경우 C: 필요한 경우 지방 '지역별' 병합 라벨 구간을 확장
    try:
        finalize_table_styles(
            ws,
            survey_type=survey_type,
            item_type=item_type,
            response_items=response_items if isinstance(response_items, list) else [],
            region_list=region_list if isinstance(region_list, list) else [],
            shrink_to_last_response=True,
        )
    except Exception:
        # 서식 오류로 변환을 중단하지 않음
        pass

    # 저장
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    wb.save(excel_path)

    return {
        "status": "success",
        "json": json_path,
        "excel": excel_path,
        "item_type": item_type,
        "item_name": item_name,
        "written_rows": written,
        "unmapped_categories": not_found,
        "normalize_report": norm_report,
    }


def process_all_json(
    json_dir: str,
    excel_dir: str,
    survey_type: str,
    org_name: str,
    extract_types: List[str] | None = None,
    progress_callback: Callable[[Dict[str, Any]], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
) -> List[Dict[str, Any]]:
    """폴더 안의 추출 JSON을 모두 Excel로 변환.

    선택 인자:
    - progress_callback: dict -> None (가능하면 호출)
    - should_cancel: () -> bool (가능하면 호출)

    GUI 진행률 표시와 소프트 스톱 버튼에서 사용한다.
    """

    os.makedirs(excel_dir, exist_ok=True)

    # 기관 설정은 한 번만 로드(권역 기본값 등에서 사용)
    org = load_org_config(org_name, project_root=_project_root())

    meta = None
    region_info = None

    def _find_by_decoded_name(target: str) -> Optional[str]:
        """파일명이 #U 디코딩 후 원하는 이름과 일치하는 경로 반환."""
        try:
            for fn in os.listdir(json_dir):
                if maybe_decode_hashu(fn) == target:
                    return os.path.join(json_dir, fn)
        except Exception:
            return None
        return None

    # 공유 메타데이터/권역 JSON이 있으면 로드('#UXXXX' 파일명도 지원)
    meta_path = _find_by_decoded_name("메타데이터.json") or os.path.join(json_dir, "메타데이터.json")
    if os.path.exists(meta_path):
        with open(meta_path, "r", encoding="utf-8") as f:
            meta = json.load(f)

    region_path = _find_by_decoded_name("권역설명.json") or os.path.join(json_dir, "권역설명.json")
    if os.path.exists(region_path):
        with open(region_path, "r", encoding="utf-8") as f:
            region_info = json.load(f)

    # 지방인데 권역설명 JSON이 없고 YAML에 고정권역이 있다면 그걸로 대체
    if survey_type == "지방" and region_info is None and getattr(org.local, "fixed_regions", []):
        region_info = {
            "권역": [{"권역명": name, "포함지역": []} for name in org.local.fixed_regions]
        }

    def _progress(current: int | None, total: int | None, message: str | None = None) -> None:
        if progress_callback is None:
            return
        try:
            progress_callback({
                "stage": "excel",
                "current": current,
                "total": total,
                "message": message,
            })
        except Exception:
            pass

    def _cancel_check() -> None:
        try:
            if should_cancel is not None and should_cancel():
                raise UserCancelled("cancelled")
        except UserCancelled:
            raise
        except Exception:
            # should_cancel 콜백 때문에 작성기가 멈추면 안 됨
            return

    results: List[Dict[str, Any]] = []

    # 결정형 진행률을 위해 작업 목록을 먼저 확정
    try:
        candidates = sorted(os.listdir(json_dir))
    except FileNotFoundError:
        raise
    except Exception:
        candidates = []

    worklist: List[str] = []
    for filename in candidates:
        if not filename.lower().endswith(".json"):
            continue
        decoded = maybe_decode_hashu(filename)
        if decoded in {"메타데이터.json", "권역설명.json", "_toc.json", "_initial_scan.json"}:
            continue
        if filename.endswith("_auto.json"):  # 자동 판별 캐시는 변환 대상 아님
            continue
        worklist.append(filename)

    total = len(worklist)
    _progress(0 if total else None, total if total else None, f"Excel 변환 준비... ({total}개)" if total else "Excel 변환 준비...")

    for idx, filename in enumerate(worklist, start=1):
        _cancel_check()
        json_path = os.path.join(json_dir, filename)
        _progress(idx, total, f"Excel 변환 중... ({idx}/{total})")

        try:
            with open(json_path, "r", encoding="utf-8") as f:
                d = json.load(f)
            item_type = (d.get("항목유형") or "ISSUE").upper()

            # GUI에서 특정 항목만 요청했다면 그 항목만 변환
            if extract_types is not None and item_type not in [t.upper() for t in extract_types]:
                continue

            item_name = d.get("항목명") or d.get("질문") or Path(filename).stem

            excel_filename = f"{item_type}_{_safe_excel_name(item_name)}.xlsx"
            excel_path = get_unique_filepath(excel_dir, excel_filename)

            res = json_to_excel(
                json_path=json_path,
                excel_path=excel_path,
                survey_type=survey_type,
                org_name=org_name,
                metadata=meta,
                region_info=region_info,
            )
            results.append(res)
        except Exception as e:
            results.append({
                "status": "failed",
                "json": json_path,
                "error": str(e),
            })

    return results


def _safe_excel_name(name: str, limit: int = 80) -> str:
    import re

    s = str(name or "").strip()
    s = re.sub(r"[\\/*?:\"<>|\n\r]", "", s)
    s = re.sub(r"\s+", " ", s)
    if not s:
        s = "item"
    return s[:limit]
