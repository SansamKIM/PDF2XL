"""API 키 없이 수동 작성한 가상 JSON을 실제 프로그램으로 Excel로 변환합니다.

JSON -> Excel 단계만 실행합니다. OpenAI API 호출이나 PDF 추출은 하지 않습니다.
필요 패키지: openpyxl, PyYAML (프로젝트 requirements.txt에 포함).

    python examples/run_example.py
    python examples/run_example.py --output examples/expected_result.xlsx
    python examples/run_example.py --check --output examples/expected_result.xlsx

기본 결과: 프로젝트의 output/examples/sample_issue.xlsx
--output 상대 경로는 현재 작업 디렉터리 기준이며, 기존 결과 파일은 덮어씁니다.
--check는 지정된 기존 결과를 검증만 하며 파일을 변경하지 않습니다.
"""

from __future__ import annotations

import argparse
from copy import copy
import json
import math
from pathlib import Path
import sys


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(__file__).with_name("sample_issue.json")
PUBLIC_AUTHOR = "pdf2xlAI example"

# 작성기의 행 매핑 함수를 재사용하지 않는 독립적인 예제 검증 기준입니다.
EXPECTED_ROWS = {
    "전체": 5,
    "서울": 6, "인천/경기": 7, "대전/세종/충청": 8, "광주/전라": 9,
    "대구/경북": 10, "부산/울산/경남": 11, "강원/제주": 12,
    "남성": 13, "여성": 14,
    "18~29세": 15, "30대": 16, "40대": 17, "50대": 18,
    "60대": 19, "70세 이상": 20,
    "진보": 21, "중도": 22, "보수": 23, "모름.무응답": 24,
    "농/임/어업": 25, "자영업": 26, "화이트칼라": 27, "블루칼라": 28,
    "전업주부": 29, "학생": 30, "기타": 31, "은퇴.무직": 32,
    "밝힐 수 없음": 33,
}
GROUP_RANGES = {
    "지역별": (6, 12), "성별": (13, 14), "연령별": (15, 20),
    "이념성향별": (21, 24), "직업별": (25, 33),
}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def validate_source(source: dict) -> None:
    """각 분류의 사례수와 가중 응답 비율이 전체와 일치하는지 확인합니다."""
    table = source["데이터"]
    headers = source["응답항목"]
    require(source["_예제정보"]["가상데이터"] is True, "가상 데이터 표시가 없습니다.")
    require(source["_예제정보"]["AI추출결과"] is False, "수동 작성 JSON 표시가 없습니다.")
    require(headers == ["찬성", "반대", "잘 모름"], "예제 응답항목이 변경되었습니다.")
    require(set(table) == set(EXPECTED_ROWS), "예제 분류가 누락되거나 추가되었습니다.")
    overall = table["전체"]
    require(overall["조사완료"] == source["메타데이터"]["표본크기"] == 1000,
            "전체 사례수는 1,000명이어야 합니다.")
    for label, values in table.items():
        require(type(values["조사완료"]) is int and values["조사완료"] > 0,
                f"{label}: 사례수는 양의 정수여야 합니다.")
        require(all(type(values[h]) in (int, float) and 0 <= values[h] <= 100
                    for h in headers), f"{label}: 비율이 0~100의 수치가 아닙니다.")
        require(math.isclose(sum(values[h] for h in headers), 100),
                f"{label}: 응답 비율의 합이 100%가 아닙니다.")
    for group, (first, last) in GROUP_RANGES.items():
        rows = [table[label] for label, row in EXPECTED_ROWS.items() if first <= row <= last]
        require(sum(row["조사완료"] for row in rows) == overall["조사완료"],
                f"{group}: 사례수 합이 전체와 다릅니다.")
        for header in headers:
            weighted = sum(row["조사완료"] * row[header] for row in rows) / overall["조사완료"]
            require(math.isclose(weighted, overall[header]),
                    f"{group}: {header}의 가중 비율이 전체와 다릅니다.")


def verify_workbook(path: Path, source: dict) -> None:
    """저장된 29개 행의 사례수·응답값과 제목·메타데이터를 원본 JSON과 대조합니다."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        fields = ["조사완료", *source["응답항목"]]
        for label, row in EXPECTED_ROWS.items():
            for col, field in enumerate(fields, start=4):
                cell = ws.cell(row, col)
                expected = source["데이터"][label][field]
                require(type(cell.value) in (int, float) and cell.value == expected,
                        f"{cell.coordinate}: {label}/{field} 값 불일치 ({cell.value!r} != {expected!r})")
        metadata = source["메타데이터"]
        expected_cells = {
            "B2": source["항목명"], "M2": source["항목명"],
            "E2": metadata["조사기관"], "G2": metadata["조사기간"],
            "H2": "가상조사", "I2": "가상표본", "J2": metadata["표본크기"],
            "K2": metadata["응답률"], "L2": metadata["표본오차"],
            "N2": 2026, "O2": 1,
            "E4": "찬성", "F4": "반대", "G4": "잘 모름",
        }
        for address, expected in expected_cells.items():
            require(ws[address].value == expected, f"{address}: 제목/헤더/메타데이터 불일치")
        require(wb.properties.creator == PUBLIC_AUTHOR and
                wb.properties.lastModifiedBy == PUBLIC_AUTHOR,
                "예제 결과의 작성자 메타데이터가 정리되지 않았습니다.")
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--output", type=Path, default=PROJECT_ROOT / "output/examples/sample_issue.xlsx",
                        help="결과 .xlsx 경로. 기본: 프로젝트/output/examples/sample_issue.xlsx")
    parser.add_argument("--check", action="store_true", help="기존 결과만 검증 (파일 생성/수정 없음)")
    args = parser.parse_args()
    output = args.output.expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("--output은 .xlsx 파일 경로여야 합니다.")
    if not args.check and output.is_relative_to(PROJECT_ROOT / "templates"):
        parser.error("원본 templates 폴더에는 결과를 저장할 수 없습니다.")

    sys.path.insert(0, str(PROJECT_ROOT))
    try:
        from openpyxl import load_workbook
        from src.excel_writer import json_to_excel

        source = json.loads(SOURCE.read_text(encoding="utf-8"))
        validate_source(source)
        if not args.check:
            result = json_to_excel(str(SOURCE), str(output), survey_type="전국", org_name="기타기관")
            require(result["status"] == "success", "Excel 변환에 실패했습니다.")
            require(result["written_rows"] == len(EXPECTED_ROWS), "일부 행이 기록되지 않았습니다.")
            require(not result["unmapped_categories"], f"매핑되지 않은 분류: {result['unmapped_categories']}")
            wb = load_workbook(output)
            try:
                # 원본 템플릿은 그대로 두고 공개 결과의 작성자 정보를 대체합니다.
                wb.properties.creator = PUBLIC_AUTHOR
                wb.properties.lastModifiedBy = PUBLIC_AUTHOR
                # 긴 예제 제목이 잘리지 않도록 결과의 제목 셀만 줄바꿈합니다.
                for address in ("B2", "M2"):
                    alignment = copy(wb.active[address].alignment)
                    alignment.wrap_text = True
                    wb.active[address].alignment = alignment
                wb.active.row_dimensions[2].height = 90
                wb.save(output)
            finally:
                wb.close()
        verify_workbook(output, source)
    except ImportError as exc:
        parser.exit(1, f"필요 패키지가 없습니다: {exc}\n설치: python -m pip install openpyxl PyYAML\n")
    except (OSError, ValueError) as exc:
        parser.exit(1, f"예제 실행 실패: {exc}\n")

    print("검증 완료: 29개 행, 사례수와 응답값 116개, 분류별 합계, 제목/헤더/메타데이터 일치")
    print("수동 작성 가상 JSON -> Excel만 실행합니다. OpenAI API/PDF 추출은 실행하지 않습니다.")
    print(f"{'검증한 파일' if args.check else '생성한 파일'}: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
